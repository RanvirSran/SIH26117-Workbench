"""
XLSX generation.

Receives an already-constructed workbook_data dictionary and writes
a real .xlsx file to backend/generated/.
"""

import os
import re
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


OUTPUT_DIR = os.path.join(
    os.path.dirname(
        os.path.dirname(
            os.path.abspath(__file__)
        )
    ),
    "generated",
)


_INVALID_SHEET_CHARS = re.compile(
    r'[\\/*?:\[\]]'
)


def _safe_sheet_name(
    name: str,
    fallback: str
) -> str:

    cleaned = _INVALID_SHEET_CHARS.sub(
        "",
        str(name or fallback)
    )

    cleaned = cleaned[:31]

    return cleaned or fallback


def generate_xlsx(
    workbook_data: dict
) -> str:
    """
    Build an .xlsx from workbook_data and return its filepath.
    """

    if not isinstance(workbook_data, dict):
        raise ValueError(
            "workbook_data must be a dictionary."
        )

    os.makedirs(
        OUTPUT_DIR,
        exist_ok=True
    )

    requested_filename = str(
        workbook_data.get(
            "filename",
            "workbook.xlsx"
        )
    ).strip()

    sheets = workbook_data.get(
        "sheets",
        []
    )

    if not isinstance(sheets, list):
        raise ValueError(
            "sheets must be a list."
        )

    wb = Workbook()

    default_sheet = wb.active

    wb.remove(
        default_sheet
    )

    header_font = Font(
        bold=True,
        color="FFFFFF"
    )

    header_fill = PatternFill(
        start_color="4472C4",
        end_color="4472C4",
        fill_type="solid"
    )

    used_names = set()

    for index, sheet_data in enumerate(sheets):

        if not isinstance(sheet_data, dict):
            raise ValueError(
                f"Sheet {index} must be a dictionary."
            )

        sheet_name = _safe_sheet_name(
            sheet_data.get(
                "name",
                ""
            ),
            f"Sheet{index + 1}"
        )

        # Prevent duplicate Excel sheet names.
        original_name = sheet_name
        suffix = 2

        while sheet_name.lower() in used_names:

            suffix_text = f"_{suffix}"

            sheet_name = (
                original_name[
                    :31 - len(suffix_text)
                ]
                + suffix_text
            )

            suffix += 1

        used_names.add(
            sheet_name.lower()
        )

        columns = sheet_data.get(
            "columns",
            []
        )

        rows = sheet_data.get(
            "rows",
            []
        )

        if not isinstance(columns, list):
            raise ValueError(
                f"Sheet '{sheet_name}' columns must be a list."
            )

        if not columns:
            raise ValueError(
                f"Sheet '{sheet_name}' has no columns."
            )

        if not isinstance(rows, list):
            raise ValueError(
                f"Sheet '{sheet_name}' rows must be a list."
            )

        # -----------------------------------------------------------
        # Validate every row BEFORE creating the file.
        # -----------------------------------------------------------

        for row_index, row in enumerate(rows):

            if not isinstance(row, list):
                raise ValueError(
                    f"Sheet '{sheet_name}' row "
                    f"{row_index} must be a list."
                )

            if len(row) != len(columns):

                raise ValueError(
                    f"Sheet '{sheet_name}' row "
                    f"{row_index} has {len(row)} values "
                    f"but there are {len(columns)} columns."
                )

        ws = wb.create_sheet(
            title=sheet_name
        )

        # -----------------------------------------------------------
        # Header
        # -----------------------------------------------------------

        for col_idx, column in enumerate(
            columns,
            start=1
        ):

            cell = ws.cell(
                row=1,
                column=col_idx,
                value=str(column)
            )

            cell.font = header_font
            cell.fill = header_fill

        # -----------------------------------------------------------
        # Rows
        # -----------------------------------------------------------

        for row_idx, row in enumerate(
            rows,
            start=2
        ):

            for col_idx, value in enumerate(
                row,
                start=1
            ):

                ws.cell(
                    row=row_idx,
                    column=col_idx,
                    value=value
                )

        # -----------------------------------------------------------
        # Column widths
        # -----------------------------------------------------------

        for col_idx, column in enumerate(
            columns,
            start=1
        ):

            max_len = len(
                str(column)
            )

            for row in rows:

                max_len = max(
                    max_len,
                    len(
                        str(
                            row[col_idx - 1]
                        )
                    )
                )

            ws.column_dimensions[
                get_column_letter(col_idx)
            ].width = min(
                max_len + 4,
                60
            )

    # ---------------------------------------------------------------
    # Empty workbook fallback
    # ---------------------------------------------------------------

    if not sheets:

        wb.create_sheet(
            title="Sheet1"
        )

    # ---------------------------------------------------------------
    # Filename
    # ---------------------------------------------------------------

    base_name = os.path.splitext(
        requested_filename
    )[0]

    safe_base = "".join(
        c
        for c in base_name
        if c.isalnum() or c in " -_"
    ).strip()

    safe_base = (
        safe_base.replace(" ", "_")
        or "workbook"
    )

    timestamp = datetime.now().strftime(
        "%Y%m%d_%H%M%S"
    )

    filename = (
        f"{safe_base}_{timestamp}.xlsx"
    )

    filepath = os.path.join(
        OUTPUT_DIR,
        filename
    )

    wb.save(
        filepath
    )

    return filepath


if __name__ == "__main__":

    test_data = {
        "filename": "safety_requirements.xlsx",

        "sheets": [
            {
                "name": "Safety Requirements",

                "columns": [
                    "Requirement",
                    "Description",
                    "Source"
                ],

                "rows": [
                    [
                        "Ventilation",
                        "Adequate ventilation must be maintained.",
                        "Oil Mines Regulations"
                    ],
                    [
                        "Protective Equipment",
                        "Workers must use appropriate protective equipment.",
                        "Safety SOP"
                    ],
                ],
            }
        ],
    }

    path = generate_xlsx(
        test_data
    )

    print(
        f"Generated: {path}"
    )

    print(
        f"Exists on disk: "
        f"{os.path.exists(path)}"
    )

    print(
        f"File size: "
        f"{os.path.getsize(path)} bytes"
    )