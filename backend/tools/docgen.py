"""
DOCX generation.

Receives an already-constructed document_data dictionary and writes
a real .docx file to backend/generated/.
"""

import os
from datetime import datetime

from docx import Document
from docx.shared import Pt


OUTPUT_DIR = os.path.join(
    os.path.dirname(
        os.path.dirname(
            os.path.abspath(__file__)
        )
    ),
    "generated",
)


def generate_docx(document_data: dict) -> str:
    """
    Build a .docx from document_data and return its filepath.

    Expected:

    {
        "title": "...",
        "sections": [
            {
                "heading": "...",
                "content": "..."
            }
        ]
    }
    """

    if not isinstance(document_data, dict):
        raise ValueError(
            "document_data must be a dictionary."
        )

    os.makedirs(
        OUTPUT_DIR,
        exist_ok=True
    )

    title = str(
        document_data.get(
            "title",
            "Untitled Document"
        )
    ).strip()

    if not title:
        title = "Untitled Document"

    sections = document_data.get(
        "sections",
        []
    )

    if not isinstance(sections, list):
        raise ValueError(
            "sections must be a list."
        )

    doc = Document()

    doc.add_heading(
        title,
        level=0
    )

    for section in sections:

        if not isinstance(section, dict):
            raise ValueError(
                "Each section must be a dictionary."
            )

        heading = str(
            section.get(
                "heading",
                ""
            )
        ).strip()

        content = str(
            section.get(
                "content",
                ""
            )
        ).strip()

        if heading:
            doc.add_heading(
                heading,
                level=1
            )

        if content:
            paragraph = doc.add_paragraph(
                content
            )

            paragraph.style.font.size = Pt(11)

    safe_title = "".join(
        c
        for c in title
        if c.isalnum() or c in " -_"
    ).strip()

    safe_title = (
        safe_title.replace(" ", "_")
        or "document"
    )

    timestamp = datetime.now().strftime(
        "%Y%m%d_%H%M%S"
    )

    filename = (
        f"{safe_title}_{timestamp}.docx"
    )

    filepath = os.path.join(
        OUTPUT_DIR,
        filename
    )

    doc.save(filepath)

    return filepath


if __name__ == "__main__":

    test_data = {
        "title": "Notice of Disease Requirements",

        "sections": [
            {
                "heading": "Requirement",
                "content": (
                    "The owner, agent or manager must give "
                    "the required notice within the applicable "
                    "period specified by the regulations."
                ),
            },
            {
                "heading": "Procedure",
                "content": (
                    "The required particulars must be submitted "
                    "to the authorities specified by the regulation."
                ),
            },
        ],
    }

    path = generate_docx(
        test_data
    )

    print(
        f"Generated: {path}"
    )

    print(
        f"Exists on disk: "
        f"{os.path.exists(path)}"
    )

    print(
        f"File size: "
        f"{os.path.getsize(path)} bytes"
    )